from pathlib import Path
import json
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

BASE = Path('/mnt/agents/output/审稿修订_本地补分析包')
WORK = BASE / 'work_v6'
OUT = Path('/mnt/agents/output/Final_Submission_Package/02_Supplementary_Tables/Supplementary_Tables_S1-S13.xlsx')
RR = BASE / 'returned_results'
MR = RR / '09_full54_scdrs_complete/results_for_return/MR'
SC = RR / '09_full54_scdrs_complete/results_for_return/scDRS'
CON = WORK / 'consolidated_results'
ACCENT = '1F4E79'
LIGHT = 'DCE6F1'
GREY = '666666'


def read(path):
    return pd.read_csv(path)


def bool_cols(df, cols):
    out = df.copy()
    for c in cols:
        if c in out:
            out[c] = out[c].map({True:'Yes',False:'No',1:'Yes',0:'No'}).fillna(out[c])
    return out


def source_id(df, value):
    out = df.copy()
    out.insert(0, 'source_id', value)
    return out

sheets = []
def add(name, group, description, df, source):
    assert len(df) > 0, name
    sheets.append(dict(name=name, group=group, description=description, df=df, source=source))

# S1 and S2: SMR catalogs and complete second-stage test set.
smr42 = read(CON / 'SMR_legacy42_global_sensitivity.csv')
smr42 = bool_cols(smr42, ['global_Bonferroni_retained_HEIDI_0_01','global_Bonferroni_retained_HEIDI_0_05','HEIDI_0_01_to_0_05'])
add('S1_SMR_42pairs','S1','Two-stage SMR candidate catalog with second-stage global-correction status.',source_id(smr42,'eQTLGen + six disease GWAS'),'eQTLGen; disease GWAS')
smrall = read(CON / 'SMR_global_correction_2354.csv')
smrall['retained_global_Bonferroni_HEIDI_0_01'] = smrall['pass_global_Bonferroni'] & smrall['pass_HEIDI_0_01']
smrall['retained_global_Bonferroni_HEIDI_0_05'] = smrall['pass_global_Bonferroni'] & smrall['pass_HEIDI_0_05']
smrall['retained_global_BH_HEIDI_0_01'] = smrall['pass_global_BH'] & smrall['pass_HEIDI_0_01']
smrall = bool_cols(smrall,['pass_within_bonferroni','pass_HEIDI_0_01','pass_HEIDI_0_05','legacy42_selected','pass_global_Bonferroni','pass_global_BH','retained_global_Bonferroni_HEIDI_0_01','retained_global_Bonferroni_HEIDI_0_05','retained_global_BH_HEIDI_0_01'])
add('S2_SMR_2354tests','S2','All 2,354 disease-probe tests entering SMR Select with within-disease and second-stage global corrections.',source_id(smrall,'eQTLGen + six disease GWAS'),'eQTLGen; disease GWAS')
smrcounts = read(CON / 'SMR_global_correction_counts.csv')
add('S2b_SMR_rule_counts','S2','Pair and unique-probe counts under each SMR correction and HEIDI rule.',source_id(smrcounts,'Derived from S2_SMR_2354tests'),'Derived results')

# S3 and S4: forward MR and complete diagnostics.
full54 = read(MR / 'S3_full54_univariate_complete.csv')
full54 = bool_cols(full54,['Steiger_correct_direction','fdr_family_full54','fdr_family_noGCC48','fdr_family_composite36'])
add('S3_forwardMR_full54','S3','Complete 54 forward MR models with five estimators, matched confidence intervals, FDR families, and input hashes.',source_id(full54,'OCT GWAS + six disease GWAS'),'OCT; disease GWAS')
presso = read(MR / 'S4_full54_MR_PRESSO_complete.csv')
presso = bool_cols(presso,['outlier_test_conducted'])
add('S4a_MR_PRESSO_54','S4','Unified 10,000-simulation MR-PRESSO results and formal outlier decisions.',source_id(presso,'Derived from S3_forwardMR_full54'),'Derived results')
loo = read(MR / 'MR_leave_one_out_full54.csv')
loo = bool_cols(loo,['direction_matches_full'])
add('S4b_MR_LOO_2045','S4','Leave-one-out estimates for all 54 forward models.',source_id(loo,'Derived from S3_forwardMR_full54'),'Derived results')
wald = read(MR / 'MR_single_SNP_Wald_full54.csv')
wald = bool_cols(wald,['palindromic'])
add('S4c_singleSNP_2045','S4','Single-SNP Wald ratios for all forward instruments.',source_id(wald,'OCT GWAS + six disease GWAS'),'OCT; disease GWAS')
input_audit = read(MR / 'MR_input_history_reconciliation.csv')
add('S4d_input_reconciliation','S4','Reconciliation of current MR inputs with historical models; valid variants were retained.',source_id(input_audit,'MR input audit'),'Derived results')
iamdgc_audit = read(MR / 'IAMDGC_extraction_audit.csv')
add('S4e_IAMDGC_audit','S4','IAMDGC 2.0 extraction and exact historical-result reconciliation for nine models.',source_id(iamdgc_audit,'IAMDGC 2.0 European subset'),'IAMDGC 2.0')
fdrfam = read(MR / 'MR_FDR_families.csv')
add('S4f_FDR_families','S4','Definitions and counts for the forward MR multiplicity families.',source_id(fdrfam,'Derived from S3_forwardMR_full54'),'Derived results')

# S5: reverse MR with updated t-based precision.
rev = read(BASE / 'inputs/mr_reverse_fdr.csv')
revp = read(RR / '03_precision/reverseMR_precision_all54.csv')
rev = rev.merge(revp.drop(columns=['nsnp','IVW_beta','IVW_se']),on=['disease','layer'],how='left',validate='one_to_one')
rev['primary_early_AMD_family'] = rev['disease'].eq('winkler_earlyamd')
rev = bool_cols(rev,['Steiger_correct','primary_early_AMD_family','reported_in_S4','negative_p_gt_0_05'])
add('S5_reverseMR_54','S5','Reverse MR estimates and diagnostics; the nine winkler_earlyamd rows form the primary reverse family.',source_id(rev,'Six disease GWAS + OCT GWAS'),'OCT; disease GWAS')
rev9 = read(RR / '03_precision/S4_reverseMR_earlyAMD_precision.csv')
rev_unified = read(CON / 'reverseMR_earlyAMD_unified.csv')
rev9 = rev9.merge(rev_unified[['disease','layer','fdr_early9_t']],on=['disease','layer'],how='left',validate='one_to_one')
rev9 = bool_cols(rev9,['reported_in_S4','negative_p_gt_0_05'])
add('S5b_reverse_early9','S5','Precision, t-based confidence intervals, detectable-effect calculations, and the primary nine-test t-based FDR for early AMD reverse models.',source_id(rev9,'Winkler early AMD + OCT GWAS'),'OCT; Winkler early AMD')

# S6: LDSC. Heritability panel values are the values displayed in the retained Figure 2 source graphic.
ld90 = read(RR / '03_precision/LDSC_rg_precision_all90.csv')
add('S6a_LDSC_rg90','S6','All 90 LDSC genetic-correlation results with precision and detectable-effect calculations.',source_id(ld90,'LDSC summary-statistic analysis'),'OCT; disease GWAS')
ld9 = ld90[ld90['trait2'].eq('winkler_earlyamd')].copy()
add('S6b_LDSC_early9','S6','Nine layer-early AMD genetic correlations with 95% confidence intervals and power-aware interpretation.',source_id(ld9,'LDSC summary-statistic analysis'),'OCT; Winkler early AMD')
h2 = pd.DataFrame({
    'trait':['RNFL','GCC','GCL','IPL','INL','OPL','PS','RPE','CSI','early AMD'],
    'h2_displayed_in_Figure_2':[0.210,0.217,0.202,0.219,0.244,0.311,0.222,0.096,0.151,0.022],
    'provenance_note':['Point estimate transcribed from the retained Figure 2 source graphic; the independent h2 standard-error table is not present in the consolidated return package.']*10
})
add('S6c_h2_figure_source','S6','Figure 2A heritability point-estimate provenance. No standard errors are imputed.',source_id(h2,'Retained Figure 2 source graphic'),'Derived results')

# S7: MVMR and robust diagnostics.
mvmr = read(RR / '02_mvmr_ucsc_verified/S7_mvmr_complete.csv')
mvmr = bool_cols(mvmr,['conditional_F_lt_10'])
add('S7a_MVMR_all18','S7','All 18 audited multivariable MR exposure estimates with conditional F and heterogeneity statistics.',source_id(mvmr,'OCT GWAS + six disease GWAS'),'OCT; disease GWAS')
for name,desc,file in [
    ('S7b_MVMR_early3','Validated three-exposure early AMD model.','MVMR_earlyAMD_validated_model.csv'),
    ('S7c_MVMR_LOO354','Complete leave-one-out refits for the three-exposure early AMD model.','MVMR_earlyAMD_LOO.csv'),
    ('S7d_MVMR_influence','Variant-level Q contributions, leverage, residuals, and Cook distance.','MVMR_earlyAMD_influence.csv'),
    ('S7e_MVMR_median','MVMR-Median parametric-bootstrap sensitivity with normal and percentile intervals.','MVMR_earlyAMD_median_sensitivity.csv'),
    ('S7f_MVMR_Egger','MVMR-Egger sensitivity including the intercept and residual heterogeneity.','MVMR_earlyAMD_Egger_sensitivity.csv')]:
    df = read(CON / file)
    df = bool_cols(df,['direction_matches_full'])
    add(name,'S7',desc,source_id(df,'Derived from the 118-instrument early AMD MVMR matrix'),'Derived results')
sel = read(MR / 'MVMR_exposure_selection_audit.csv')
add('S7g_exposure_selection','S7','Outcome-specific exposure-selection audit for the full54 analysis.',source_id(sel,'Derived from S3_forwardMR_full54'),'Derived results')

# S8: sample-overlap scenarios.
overlap = read(CON / 'earlyAMD_sample_overlap_bias_scenarios.csv')
overlap = bool_cols(overlap,['shifted_ci_excludes_zero'])
add('S8_overlap_scenarios','S8','All 360 explicit sample-overlap bias scenarios using worst-toward-null shifts.',source_id(overlap,'Burgess-Davies-Thompson scenario approximation'),'Derived results')
ovsum = overlap.groupby(['layer','strength_parameter','absolute_observational_logOR_assumption'],as_index=False).agg(
    scenarios=('relative_overlap_fraction','size'),
    max_relative_overlap=('relative_overlap_fraction','max'),
    max_approximate_absolute_bias=('approximate_absolute_bias','max'),
    all_shifted_CIs_exclude_zero=('shifted_ci_excludes_zero',lambda x: 'Yes' if (x=='Yes').all() else 'No'))
add('S8b_overlap_summary','S8','Scenario-level summary by layer, instrument-strength choice, and observational-association assumption.',source_id(ovsum,'Derived from S8_overlap_scenarios'),'Derived results')

# S9: colocalization.
col400 = read(BASE / 'inputs/coloc_regions.csv')
col400 = col400[col400['disease'].eq('winkler_earlyamd')].copy()
add('S9a_coloc_early400','S9','All 400 early AMD layer-window colocalization tests under the specified single-causal-variant model.',source_id(col400,'OCT GWAS + Winkler early AMD'),'OCT; Winkler early AMD')
col13 = read(RR / '04_smr_coloc/S5a_earlyAMD_H4_gt_0.9_region_audit.csv')
add('S9b_coloc_H4_13','S9','Thirteen PP.H4>0.9 tests with distinct-window and connected-interval identifiers.',source_id(col13,'Derived from S9a_coloc_early400'),'Derived results')
conn = read(RR / '04_smr_coloc/earlyAMD_connected_intervals.csv')
add('S9c_connected_intervals','S9','Eight connected intervals formed by cross-layer overlap of high-support windows.',source_id(conn,'Derived from S9b_coloc_H4_13'),'Derived results')
colspec = read(RR / '04_smr_coloc/S5b_specificity_H4_gt_0.5.csv')
add('S9d_specificity_H4','S9','Specificity-panel windows with PP.H4>0.5.',source_id(colspec,'OCT GWAS + specificity-panel GWAS'),'OCT; specificity GWAS')
speccounts = read(RR / '04_smr_coloc/specificity_coloc_counts_by_disease.csv')
add('S9e_specificity_counts','S9','Descriptive specificity-panel colocalization counts by disease.',source_id(speccounts,'Derived from specificity colocalization results'),'Derived results')

# S10: scDRS and atlas.
scmain = read(SC / 'main/scdrs_group_summary.csv')
add('S10a_scDRS_main66','S10','Primary 66-test scDRS analysis with within-trait and global66 FDR.',source_id(scmain,'Adult human retinal atlas + MAGMA gene sets'),'CELLxGENE; GWAS')
scpar = read(SC / 'sensitivity/scdrs_parameter_sensitivity_summary.csv')
scpar = bool_cols(scpar,['direction_concordant'])
add('S10b_scDRS_parameters','S10','Parameter sensitivity using top 500, top 2,000, and 10 kb window designs.',source_id(scpar,'scDRS sensitivity analysis'),'Derived results')
sclodo = read(SC / 'lodo/scdrs_LODO_summary.csv')
sclodo = bool_cols(sclodo,['scores_recomputed_after_donor_omission','direction_concordant'])
add('S10c_scDRS_LODO','S10','Leave-one-donor-out analyses with preprocessing and scoring repeated after each omission.',source_id(sclodo,'scDRS leave-one-donor-out analysis'),'Derived results')
atlas = json.load(open(RR / '09_full54_scdrs_complete/results_for_return/atlas/atlas_build_manifest.json'))
annot = pd.DataFrame([
    ['NN','glial cell',1473],['HC','retina horizontal cell',601],['AC','amacrine cell',4623],
    ['BC','retinal bipolar neuron',4339],['Rod','retinal rod cell',8012],['Cone','retinal cone cell',951]
],columns=['author_cell_type','cell_type_ontology_label','n_cell'])
annot['annotation_relationship'] = 'One-to-one; counted once in the primary 66-test family'
add('S10d_atlas_annotation','S10','Author and ontology labels, exact cell counts, and one-to-one annotation mapping.',source_id(annot,'Adult human retinal atlas'),'CELLxGENE')
donors = pd.DataFrame({'donor_id':list(atlas['donor_counts'].keys()),'n_cell':list(atlas['donor_counts'].values())})
donors['n_donors'] = 6
add('S10e_atlas_donors','S10','Donor-level cell counts in the verified 19,999-cell mini atlas.',source_id(donors,'Adult human retinal atlas'),'CELLxGENE')

# S11: Visium sensitivity and descriptive expression.
viscorr = read(RR / '05_visium_sensitivity/S6a_Visium_correlation_sensitivity.csv')
add('S11a_Visium_correlations','S11','Raw, technically adjusted, and matched-random empirical spatial-correlation results.',source_id(viscorr,'GSE234047'),'GSE234047')
vispanel = read(RR / '05_visium_sensitivity/S6d_Visium_16gene_panel_audit.csv')
add('S11b_Visium_panel80','S11','Sixteen candidates by five sections with valid aliases, detection, zero-variance, and missing status.',source_id(vispanel,'GSE234047 + annotation audit'),'GSE234047')
oldsup = Path('/mnt/agents/output/手稿/补充表_S1-S6.xlsx')
visde = pd.read_excel(oldsup,sheet_name='S6b_Visium_DE')
add('S11c_Visium_DEScriptive','S11','Descriptive spot-level expression differences between CNV and control spots; donor and disease are confounded.',source_id(visde,'GSE234047'),'GSE234047')
vissec = pd.read_excel(oldsup,sheet_name='S6c_Visium_section_sensitivity')
vissec = vissec.drop(columns=[c for c in ['Source Name','Source URL'] if c in vissec])
add('S11d_Visium_sections','S11','Descriptive section-level expression relative to the unweighted control-section mean.',source_id(vissec,'GSE234047'),'GSE234047')
vismatch = read(RR / '05_visium_sensitivity/Visium_random_gene_matching_audit.csv')
add('S11e_random_matching','S11','Audit of detection-rate and expression matching for the random gene-set null.',source_id(vismatch,'GSE234047 random-null audit'),'GSE234047')
visalias = read(RR / '05_visium_sensitivity/Visium_ENSG00000258311_annotation_audit.csv')
add('S11f_ENSG_alias_audit','S11','ENSG00000258311 alias audit and exclusion of the erroneous RP11-762I7.5 assignment.',source_id(visalias,'Ensembl GRCh37 + GSE234047'),'Ensembl; GSE234047')

# S12: protein analyses.
med = read(RR / '08_pqtl_grch37_reanalysis/S8_mediation_coordinate_sensitivity.csv')
add('S12a_protein_mediation','S12','Exploratory protein-mediation and GRCh37 coordinate sensitivity, including the non-estimable PDGFB chain.',source_id(med,'OCT GWAS + INTERVAL + disease GWAS'),'INTERVAL; GWAS')
stepa = json.load(open(RR / '08_pqtl_grch37_reanalysis/stepA_results.json'))
stepb = json.load(open(RR / '08_pqtl_grch37_reanalysis/stepB_results.json'))
steps = []
for chain,x in stepa.items():
    steps.append(['step A',chain,x.get('layer'),x.get('protein'),None,x.get('nsnp'),x.get('b_LA'),x.get('se_LA'),x.get('p'),x.get('Q'),None])
for chain,x in stepb.items():
    steps.append(['step B',chain,None,x.get('protein'),x.get('disease'),x.get('nsnp'),x.get('b_PD'),x.get('se_PD'),x.get('p'),x.get('Q'),';'.join(x.get('instruments',[]))])
steps = pd.DataFrame(steps,columns=['step','chain','layer','protein','disease','n_instruments','beta','se','p','Q','instruments'])
add('S12b_protein_steps','S12','Step A and step B protein analyses, including ERBB3 and the zero-instrument PDGFB step.',source_id(steps,'OCT GWAS + INTERVAL + disease GWAS'),'INTERVAL; GWAS')
pqcoord = read(RR / '08_pqtl_grch37_reanalysis/S8_pQTL_coordinate_audit.csv')
add('S12c_pQTL_coordinates','S12','Protein-instrument coordinate audit.',source_id(pqcoord,'INTERVAL GRCh37 coordinate audit'),'INTERVAL')

# S13: provenance, sample sizes, software, parameters, and QA.
prov = read(BASE / 'config/data_provenance_audit.csv').drop(columns=['remaining_local_check'])
prov = prov.rename(columns={'source_url':'Source URL','resource':'Source Name'})
prov.loc[prov['Source Name'].eq('IAMDGC late AMD'),'Source Name'] = 'IAMDGC 2.0 late AMD European subset'
add('S13a_data_sources','S13','Public data resources, verified identifiers or versions, and source URLs.',source_id(prov,'Public resource audit'),'Public sources')
octn = read(RR / '06_sample_sizes/OCT_GWAS_sample_sizes_by_layer.csv').drop(columns=['source_file'])
add('S13b_OCT_sample_sizes','S13','Full streaming audit of variant-specific OCT GWAS sample sizes.',source_id(octn,'Nine processed OCT GWAS files'),'OCT GWAS')
soft = json.load(open(RR / '07_environment/software_environment.json'))
software = [
    ['Python','3.14.4','Local returned analysis environment','software_environment.json'],
    ['Python','3.12','Sandbox reconciliation and figure environment','Consolidated analysis runtime'],
    ['PLINK','v1.90b7.1 64-bit (18 Oct 2023)','Completed','software_environment.json'],
    ['SMR','v1.3.2 Linux, build 19 Jun 2025','Completed local reproduction','software_environment.json'],
    ['MAGMA','v1.10','Completed full54/scDRS pipeline','MAGMA manifest'],
    ['scDRS','v1.0.2','Completed full54/scDRS pipeline','scDRS manifests'],
    ['Rscript','not found','Not used for the returned full54/scDRS pipeline','software_environment.json'],
    ['MR implementation','Project Python mrtools','Actual MR implementation; TwoSampleMR was not used','software_environment.json'],
    ['MR-PRESSO implementation','Independent Python numerical implementation','10,000 simulations; not the R package','Final MR-PRESSO outputs'],
]
software = pd.DataFrame(software,columns=['software_or_component','version_or_status','use_status','evidence_file'])
add('S13c_software','S13','Software actually used or audited, with no unsupported TwoSampleMR claim.',source_id(software,'Software environment and analysis manifests'),'Environment')
params = [
    ['Forward MR','Primary family','54 tests; IVW t df=nsnp-1; MR-Egger t df=nsnp-2'],
    ['Forward MR','Median and modes','10,000 nonparametric bootstrap iterations; seeds 1, 2, and 3'],
    ['Mode estimators','Kernel and grid','Hartwig normal kernel; Silverman bandwidth; 2,000-point grid'],
    ['MR-PRESSO','Null distribution','10,000 simulations; model-specific seeds in MR_PRESSO_seed_map.json'],
    ['MR-PRESSO','Outlier rule','Global P<0.05 required; strict-exceedance raw P; Bonferroni across instruments'],
    ['MVMR','Primary inference','t df=nsnp-number of exposures; conditional F with gencov=0'],
    ['MVMR-Median','Bootstrap','10,000 parametric iterations; seed 20261026; original outcome-SE weights'],
    ['Colocalization','Priors','p1=1e-4; p2=1e-4; p12=5e-6; W1=(0.15*sdY)^2; W2=0.15^2'],
    ['Colocalization','Regions','Lead position plus and minus 500 kb; overlaps merged within layer'],
    ['scDRS','Main design','Top 1,000 genes; 1,000 matched control sets; score random_seed=0'],
    ['scDRS','MAGMA window','35 kb upstream and 10 kb downstream; sensitivity 10 kb and 10 kb'],
    ['scDRS','Preprocessing','Raw counts; size-factor to 1e4 and log1p; min 250 genes per cell; min 50 cells per gene'],
    ['scDRS','Sensitivity','Top 500 and top 2,000 genes; six leave-one-donor-out analyses'],
    ['Visium','Random null','1,000 matched random gene sets per section; add-one empirical P; seeds 20261024 to 20261028'],
    ['Visium','Technical adjustment','Residualized ranks for log1p(nCount) and log1p(nFeature)'],
]
params = pd.DataFrame(params,columns=['analysis','parameter','final_value_or_rule'])
add('S13d_final_parameters','S13','Unified final parameters that supersede historical configuration values.',source_id(params,'Final analysis manifests'),'Parameters')
qa = [
    ['Returned complete package','Independent acceptance','26/26 checks passed'],
    ['Returned file inventory','SHA256 and byte counts','334 files passed'],
    ['Consolidated numerical analysis','Independent QA','149/149 checks passed'],
    ['IAMDGC models','Current-source exact matches','9/9'],
    ['Other historical MR models','Accounted for','29 exact matches plus 16 explained single-variant historical omissions'],
    ['Early AMD MVMR','Independent reconciliation','18/18 exposure estimates matched; tolerance 1e-6'],
]
qa = pd.DataFrame(qa,columns=['component','check','result'])
add('S13e_analysis_QA','S13','Completed numerical and file-integrity checks for the returned and consolidated results.',source_id(qa,'QA reports in the reproducibility archive'),'QA')

# Create workbook.
wb = Workbook()
wb.remove(wb.active)
readme = wb.create_sheet('README')
readme.sheet_view.showGridLines = False
readme['A1'] = 'Supplementary Tables S1 to S13'
readme['A1'].font = Font(name='Arial',size=18,bold=True,color=ACCENT)
readme.merge_cells('A1:F1')
notes = [
    'This workbook accompanies the final manuscript and contains complete result tables rather than a reviewer-response document.',
    'Evidence tiers are kept distinct: primary statistical models, internal sensitivity analyses, and descriptive molecular or tissue context.',
    'SMR global corrections apply to the second-stage tests and do not calibrate the first-stage disease-GWAS-guided region selection.',
    'The RPE multivariable finding is single-variant-dependent and estimator-sensitive. It is not presented as an established protective mechanism.',
    'Visium analyses are descriptive because disease status is completely confounded with donor identity and matched random gene-set tests were null.',
    'The GCL-TNXB mediation pipeline retains its historical 32-instrument input and is separate from the 33-instrument full54 forward model.',
    'S6c records only the heritability values displayed in the retained Figure 2 graphic. It does not impute missing h2 standard errors.',
    'All tables use three-line formatting: thick top and bottom rules, a thin header rule, and no vertical or internal horizontal rules.',
    'Source IDs map to S13a_data_sources and to the reproducibility archive. No public code repository or archived DOI is claimed.',
]
for i,text in enumerate(notes,3):
    readme.cell(i,1,u'• '+text)
    readme.merge_cells(start_row=i,start_column=1,end_row=i,end_column=6)
    readme.cell(i,1).alignment = Alignment(wrap_text=True,vertical='top')
    readme.cell(i,1).font = Font(name='Arial',size=10)
    readme.row_dimensions[i].height = 32
readme.column_dimensions['A'].width = 22
for c in 'BCDEF': readme.column_dimensions[c].width = 18

index = wb.create_sheet('Table_Index')
index.sheet_view.showGridLines = False
headers = ['Sheet','Group','Description','Records','Source']
for j,h in enumerate(headers,1): index.cell(1,j,h)

thin = Side(style='thin',color='000000')
thick = Side(style='medium',color='000000')
none = Side(style=None)

def style_table(ws, nrows, ncols, freeze='A2'):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f'A1:{get_column_letter(ncols)}{nrows+1}'
    for cell in ws[1]:
        cell.font = Font(name='Arial',size=9,bold=True,color='FFFFFF')
        cell.fill = PatternFill('solid',fgColor=ACCENT)
        cell.alignment = Alignment(vertical='center',wrap_text=True)
        cell.border = Border(top=thick,bottom=thin)
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2,max_row=nrows+1,max_col=ncols):
        for cell in row:
            cell.font = Font(name='Arial',size=9)
            cell.alignment = Alignment(vertical='top',wrap_text=False)
            cell.border = Border(bottom=none)
    for cell in ws[nrows+1]:
        cell.border = Border(bottom=thick)
    # Widths use sampled content and are capped to keep sheets navigable.
    for j in range(1,ncols+1):
        vals = [ws.cell(1,j).value] + [ws.cell(i,j).value for i in range(2,min(nrows+1,202))]
        length = max((len(str(v)) for v in vals if v is not None),default=8)
        ws.column_dimensions[get_column_letter(j)].width = min(max(length+2,11),42)
    ws.sheet_view.zoomScale = 85

for i,item in enumerate(sheets,2):
    name=item['name']; df=item['df']
    index.cell(i,1,name); index.cell(i,2,item['group']); index.cell(i,3,item['description'])
    index.cell(i,4,f"=COUNTA('{name}'!A2:A{len(df)+1})")
    index.cell(i,5,item['source'])
    ws = wb.create_sheet(name)
    # Replace NaN with None and numpy scalars with native Python values.
    clean = df.astype(object).where(pd.notna(df),None)
    for j,col in enumerate(clean.columns,1): ws.cell(1,j,str(col))
    for r,row in enumerate(clean.itertuples(index=False,name=None),2):
        for j,val in enumerate(row,1):
            if isinstance(val,np.generic): val = val.item()
            ws.cell(r,j,val)
    style_table(ws,len(df),len(df.columns))
    # Numeric display precision without converting values to text.
    for j,col in enumerate(clean.columns,1):
        lc = str(col).lower()
        if any(x in lc for x in ['beta','se_','_se','ci95','ci_','rg','mcz','h2','proportion','prop_','correlation','spearman']):
            for r in range(2,len(df)+2): ws.cell(r,j).number_format = '0.000000'
        if lc in ['p','p_t','ivw_p','p_smr','p_heidi','global_p','q_full54','assoc_mcp','assoc_fdr_global66'] or lc.endswith('_p') or lc.startswith('p_') or lc.startswith('q_'):
            for r in range(2,len(df)+2): ws.cell(r,j).number_format = '0.000E+00'

style_table(index,len(sheets),5)
index.column_dimensions['A'].width = 28
index.column_dimensions['B'].width = 10
index.column_dimensions['C'].width = 85
index.column_dimensions['D'].width = 12
index.column_dimensions['E'].width = 35
for row in index.iter_rows(min_row=2,max_row=len(sheets)+1):
    row[2].alignment = Alignment(wrap_text=True,vertical='top')
    row[4].alignment = Alignment(wrap_text=True,vertical='top')
    index.row_dimensions[row[0].row].height = 30

# Workbook metadata and calculation settings.
wb.properties.title = 'Supplementary Tables S1 to S13'
wb.properties.subject = 'Retinal layer architecture and macular disease'
wb.properties.creator = 'Study authors'
try:
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = 'auto'
except Exception:
    pass
OUT.parent.mkdir(parents=True,exist_ok=True)
wb.save(OUT)
print(f'Wrote {OUT} with {len(sheets)} data sheets plus README and Table_Index.')
