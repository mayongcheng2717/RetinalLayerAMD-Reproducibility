from pathlib import Path
import zipfile, tempfile, os, shutil, re, subprocess, sys, json, hashlib
from openpyxl import load_workbook

OUT = Path('/mnt/agents/output/Final_Submission_Package/02_Supplementary_Tables/Supplementary_Tables_S1-S13.xlsx')
GATE = Path('/app/.agents/plugins/kimi-excel/skills/kimi-excel/scripts/xlsx_tools.py')

def view_xml(freeze, zoom, selected):
    selected_attr = ' tabSelected="true"' if selected else ''
    if freeze:
        return (f'<sheetViews><sheetView showGridLines="false" zoomScale="{zoom}" zoomScaleNormal="{zoom}" workbookViewId="0"{selected_attr}>'
                f'<pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/>'
                f'<selection pane="bottomLeft" activeCell="A2" sqref="A2"/></sheetView></sheetViews>')
    return f'<sheetViews><sheetView showGridLines="false" zoomScale="{zoom}" zoomScaleNormal="{zoom}" workbookViewId="0"{selected_attr}><selection activeCell="A1" sqref="A1"/></sheetView></sheetViews>'

wb = load_workbook(OUT, data_only=False)
fd, temp_name = tempfile.mkstemp(suffix='.xlsx')
os.close(fd)
with zipfile.ZipFile(OUT, 'r') as source, zipfile.ZipFile(temp_name, 'w', zipfile.ZIP_DEFLATED) as target:
    for item in source.infolist():
        data = source.read(item.filename)
        match = re.fullmatch(r'xl/worksheets/sheet(\d+)\.xml', item.filename)
        if match:
            idx = int(match.group(1))
            if 1 <= idx <= len(wb.worksheets):
                text = data.decode('utf-8')
                text, count = re.subn(r'<sheetViews>.*?</sheetViews>', view_xml(idx != 1, 90 if idx == 1 else 85, idx == 1), text, count=1)
                if count != 1:
                    raise RuntimeError(f'Could not patch sheet view: {item.filename}')
                data = text.encode('utf-8')
        target.writestr(item, data)
shutil.move(temp_name, OUT)

result = subprocess.run([sys.executable, str(GATE), 'verify', str(OUT)], capture_output=True, text=True, timeout=300)
print(result.stdout.strip())
if result.returncode != 0:
    raise RuntimeError(result.stderr)

wbf = load_workbook(OUT, data_only=False)
wbv = load_workbook(OUT, data_only=True)
issues = []
for i, ws in enumerate(wbf.worksheets, 1):
    if ws.sheet_view.showGridLines is not False:
        issues.append(f'gridlines: {ws.title}')
    if i > 1 and ws.freeze_panes != 'A2':
        issues.append(f'freeze: {ws.title}={ws.freeze_panes}')
    if i > 1:
        for cell in ws[1]:
            if cell.border.top.style != 'medium' or cell.border.bottom.style != 'thin':
                issues.append(f'header border: {ws.title}!{cell.coordinate}')
        for cell in ws[ws.max_row]:
            if cell.border.bottom.style != 'medium':
                issues.append(f'last-row border: {ws.title}!{cell.coordinate}')
# Independent formula-cache check against actual sheet row counts.
for i in range(2, wbv['Table_Index'].max_row + 1):
    sheet = wbv['Table_Index'].cell(i, 1).value
    expected = wbf[sheet].max_row - 1
    actual = wbv['Table_Index'].cell(i, 4).value
    if actual != expected:
        issues.append(f'formula cache: {sheet} {actual} != {expected}')
if issues:
    raise RuntimeError(json.dumps(issues, ensure_ascii=False))
print(json.dumps({'status':'PASS','sheets':len(wbf.sheetnames),'data_sheets':len(wbf.sheetnames)-2,'sha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),'bytes':OUT.stat().st_size},indent=2))
